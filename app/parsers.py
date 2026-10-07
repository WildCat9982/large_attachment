"""Turn an uploaded file into either text segments or SQLite tables.

Prose (pdf, docx, txt, html, free-form json) is read; data (csv, xlsx,
record-shaped json) is loaded into a per-document SQLite file and queried.
"""
import csv
import json
import re
import sqlite3

import pandas as pd


class NoTables(Exception):
    """The file holds no record-shaped data; treat it as text instead."""


# --- text --------------------------------------------------------------------

def _read_text(path):
    with open(path, "rb") as f:
        raw = f.read()
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        return raw.decode("latin-1")


def extract_text(path, ext):
    """Yield (loc, text) segments in document order."""
    if ext == "pdf":
        from pypdf import PdfReader
        reader = PdfReader(path)
        for i, page in enumerate(reader.pages, start=1):
            try:
                text = page.extract_text() or ""
            except Exception:
                text = ""  # one unreadable page should not sink the document
            yield f"p.{i}", text
    elif ext == "docx":
        import docx
        from docx.table import Table
        parts = []
        for block in docx.Document(path).iter_inner_content():
            if isinstance(block, Table):
                for row in block.rows:
                    parts.append(" | ".join(cell.text.strip() for cell in row.cells))
            else:
                parts.append(block.text)
        yield "", "\n".join(parts)
    elif ext in ("html", "htm"):
        from bs4 import BeautifulSoup
        soup = BeautifulSoup(_read_text(path), "lxml")
        for tag in soup(["script", "style", "noscript", "template", "svg"]):
            tag.decompose()
        yield "", soup.get_text("\n")
    elif ext == "json":
        yield "", json.dumps(json.loads(_read_text(path)), indent=1, ensure_ascii=False)
    else:
        yield "", _read_text(path)


def chunk_text(text, size):
    """Split into pieces of about `size` chars, breaking at line or word ends."""
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    i, n = 0, len(text)
    while i < n:
        j = min(i + size, n)
        if j < n:
            cut = text.rfind("\n", i + size // 2, j)
            if cut == -1:
                cut = text.rfind(" ", i + size // 2, j)
            if cut != -1:
                j = cut
        piece = text[i:j].strip()
        if piece:
            yield piece
        i = j


# --- tables ------------------------------------------------------------------

def _ident(name, used, fallback):
    s = re.sub(r"\W+", "_", str(name if name is not None else "")).strip("_").lower() or fallback
    if s[0].isdigit():
        s = f"{fallback}_{s}"
    base, k = s, 2
    while s in used:
        s, k = f"{base}_{k}", k + 1
    used.add(s)
    return s


def _idents(names, fallback="col"):
    used = set()
    return [_ident(n, used, fallback) for n in names]


def _load_csv(path, conn):
    with open(path, "rb") as f:
        head = f.read(65536)
    try:
        head.decode("utf-8")
        encoding = "utf-8-sig"
    except UnicodeDecodeError as e:
        # an error at the very end is just a character cut by the 64KB read
        encoding = "utf-8-sig" if e.start > len(head) - 4 else "latin-1"
    try:
        sep = csv.Sniffer().sniff(head.decode(encoding, errors="replace"), delimiters=",;\t|").delimiter
    except csv.Error:
        sep = ","
    cols = None
    for df in pd.read_csv(path, sep=sep, encoding=encoding, chunksize=50_000, low_memory=False):
        if cols is None:
            cols = _idents(df.columns)
        df.columns = cols
        df.to_sql("data", conn, if_exists="append", index=False)
    if cols is None:
        raise ValueError("The CSV file has no rows.")
    return ["data"]


def _cell(v):
    return v if isinstance(v, (int, float, str, type(None))) else str(v)


def _load_xlsx(path, conn):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    tables, used = [], set()
    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        header = next((r for r in rows if r and any(c is not None for c in r)), None)
        if header is None:
            continue
        cols = _idents(header)
        table = _ident(ws.title, used, "sheet")
        conn.execute(f'CREATE TABLE "{table}" ({", ".join(chr(34) + c + chr(34) for c in cols)})')
        insert = f'INSERT INTO "{table}" VALUES ({",".join("?" * len(cols))})'
        batch = []
        for r in rows:
            if not r or all(c is None for c in r):
                continue
            vals = [_cell(v) for v in r[:len(cols)]]
            batch.append(vals + [None] * (len(cols) - len(vals)))
            if len(batch) >= 10_000:
                conn.executemany(insert, batch)
                batch = []
        if batch:
            conn.executemany(insert, batch)
        tables.append(table)
    wb.close()
    if not tables:
        raise ValueError("The workbook has no data.")
    return tables


def _is_records(v):
    return isinstance(v, list) and len(v) > 0 and all(isinstance(x, dict) for x in v[:100])


def _load_json(path, conn):
    data = json.loads(_read_text(path))
    if _is_records(data):
        sets = {"data": data}
    elif isinstance(data, dict):
        sets = {k: v for k, v in data.items() if _is_records(v)}
    else:
        sets = {}
    if not sets:
        raise NoTables()
    tables, used = [], set()
    for name, records in sets.items():
        table = _ident(name, used, "t")
        for start in range(0, len(records), 20_000):
            df = pd.json_normalize(records[start:start + 20_000], max_level=2)
            df = df.astype(object).where(df.notna(), None)
            df = df.map(lambda v: json.dumps(v, ensure_ascii=False) if isinstance(v, (list, dict)) else v)
            df.columns = [re.sub(r"\W+", "_", str(c)).strip("_").lower() or "col" for c in df.columns]
            if start == 0:
                df.to_sql(table, conn, index=False)
                known = set(df.columns)
            else:
                for c in df.columns:
                    if c not in known:
                        conn.execute(f'ALTER TABLE "{table}" ADD COLUMN "{c}"')
                        known.add(c)
                df.to_sql(table, conn, if_exists="append", index=False)
        tables.append(table)
    return tables


def load_tables(path, ext, out_db):
    """Load a data file into SQLite at out_db. Returns the table names."""
    conn = sqlite3.connect(out_db)
    try:
        with conn:
            if ext == "csv":
                return _load_csv(path, conn)
            if ext == "xlsx":
                return _load_xlsx(path, conn)
            if ext == "json":
                return _load_json(path, conn)
        raise NoTables()
    finally:
        conn.close()


# --- describing and querying tables -----------------------------------------

def _short(v, n=60):
    s = str(v)
    return s if len(s) <= n else s[: n - 1] + "…"


def describe_tables(db_path, tables, sample_rows=3):
    """Schema text given to the model: tables, columns, types, sample rows."""
    conn = sqlite3.connect(f"file:{db_path}?mode=ro", uri=True)
    out = []
    try:
        for t in tables:
            n = conn.execute(f'SELECT COUNT(*) FROM "{t}"').fetchone()[0]
            cols = [r[1] for r in conn.execute(f'PRAGMA table_info("{t}")')]
            sample = conn.execute(f'SELECT * FROM "{t}" LIMIT {sample_rows}').fetchall()
            types = []
            for i, c in enumerate(cols):
                kinds = {type(row[i]).__name__ for row in sample if row[i] is not None}
                if len(kinds) == 1:
                    types.append({"int": "integer", "float": "real", "str": "text"}.get(kinds.pop(), "other"))
                else:
                    types.append("mixed" if kinds else "unknown")
            out.append(f'Table "{t}" ({n:,} rows)')
            out.append("Columns: " + ", ".join(f'"{c}" {ty}' for c, ty in zip(cols, types)))
            for row in sample:
                out.append("  sample: " + " | ".join(_short(v, 40) for v in row))
            out.append("")
    finally:
        conn.close()
    return "\n".join(out).strip()


def profile_tables(db_path, tables, max_cols=40):
    """Per-column statistics used to summarise a data file."""
    conn = sqlite3.connect(f"file:{db_path}?mode=ro", uri=True)
    out = []
    try:
        for t in tables:
            n = conn.execute(f'SELECT COUNT(*) FROM "{t}"').fetchone()[0]
            cols = [r[1] for r in conn.execute(f'PRAGMA table_info("{t}")')]
            out.append(f'Table "{t}": {n:,} rows, {len(cols)} columns')
            for c in cols[:max_cols]:
                q = f'"{c}"'
                filled, distinct, lo, hi, avg = conn.execute(
                    f"SELECT COUNT({q}), COUNT(DISTINCT {q}), MIN({q}), MAX({q}), "
                    f"AVG(CASE WHEN typeof({q}) IN ('integer','real') THEN {q} END) FROM \"{t}\""
                ).fetchone()
                line = f"- {c}: {filled:,} filled, {distinct:,} distinct, min={_short(lo)}, max={_short(hi)}"
                if avg is not None:
                    line += f", mean={avg:.4g}"
                if 0 < distinct <= 30:
                    top = conn.execute(
                        f'SELECT {q}, COUNT(*) FROM "{t}" WHERE {q} IS NOT NULL '
                        f"GROUP BY {q} ORDER BY COUNT(*) DESC LIMIT 5").fetchall()
                    line += "; top: " + ", ".join(f"{_short(v, 30)} ({k:,})" for v, k in top)
                out.append(line)
            if len(cols) > max_cols:
                out.append(f"- ({len(cols) - max_cols} more columns not profiled)")
            out.append("")
    finally:
        conn.close()
    return "\n".join(out).strip()


def run_sql(db_path, sql, max_rows, timeout_s):
    """Run one read-only SELECT with a row cap and a time limit."""
    import time
    stmt = sql.strip().rstrip(";").strip()
    if ";" in stmt or not re.match(r"(?is)^(select|with)\b", stmt):
        raise ValueError("Only a single SELECT statement is allowed.")
    conn = sqlite3.connect(f"file:{db_path}?mode=ro", uri=True)  # read-only at the file level
    deadline = time.time() + timeout_s
    conn.set_progress_handler(lambda: 1 if time.time() > deadline else 0, 10_000)
    try:
        cur = conn.execute(stmt)
        cols = [d[0] for d in cur.description]
        rows = cur.fetchmany(max_rows + 1)
    except sqlite3.OperationalError as e:
        if "interrupted" in str(e).lower():
            raise TimeoutError(f"Query exceeded {timeout_s}s; make it more selective.") from e
        raise
    finally:
        conn.close()
    return cols, rows[:max_rows], len(rows) > max_rows
